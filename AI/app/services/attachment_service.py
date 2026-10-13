import io
import os
import re
import shutil
from functools import lru_cache
from pathlib import Path, PurePath

from app.core.config import get_settings
from app.core.exceptions import UnsupportedAttachmentError

MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 50_000

# OCR for scanned PDF pages
MIN_PAGE_TEXT_CHARS = 20
MAX_OCR_PAGES = 30
OCR_DPI = 300

TEXT_EXTENSIONS = {".txt", ".csv", ".tsv", ".md", ".json", ".xml", ".log", ".yaml", ".yml"}
HTML_EXTENSIONS = {".html", ".htm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

MIME_TO_EXTENSION = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "text/plain": ".txt",
    "text/csv": ".csv",
    "text/html": ".html",
}


def is_image(filename: str, mime_type: str) -> bool:
    return _extension(filename, mime_type) in IMAGE_EXTENSIONS


def extract_text(filename: str, mime_type: str, data: bytes) -> str:
    """Return the readable text of an attachment.

    Raises UnsupportedAttachmentError when the file type cannot be read.
    """
    extension = _extension(filename, mime_type)
    if extension == ".pdf":
        text = _pdf(data)
    elif extension == ".docx":
        text = _docx(data)
    elif extension in (".xlsx", ".xlsm"):
        text = _xlsx(data)
    elif extension in TEXT_EXTENSIONS:
        text = _decode(data)
    elif extension in HTML_EXTENSIONS:
        text = html_to_text(_decode(data))
    elif extension in IMAGE_EXTENSIONS:
        text = _image(data)
    else:
        raise UnsupportedAttachmentError(
            f"Unsupported file type '{extension or mime_type}'"
        )
    return text.strip()


def html_to_text(markup: str) -> str:
    import html

    markup = re.sub(r"(?is)<(script|style|head)[^>]*>.*?</\1>", " ", markup)
    markup = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h[1-6])>", "\n", markup)
    text = html.unescape(re.sub(r"<[^>]+>", " ", markup))
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _extension(filename: str, mime_type: str) -> str:
    extension = PurePath(filename).suffix.lower()
    if mime_type.startswith("image/") and extension not in IMAGE_EXTENSIONS:
        return ".png"
    return extension or MIME_TO_EXTENSION.get(mime_type, "")


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes) -> str:
    import pdfplumber

    pages = []
    ocr_pages = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            # Scanned pages have no (or almost no) text layer: OCR the rendering.
            if len(text.strip()) < MIN_PAGE_TEXT_CHARS and ocr_pages < MAX_OCR_PAGES:
                image = page.to_image(resolution=OCR_DPI).original
                text = _ocr(image)
                ocr_pages += 1
            pages.append(text)
            if sum(len(p) for p in pages) > MAX_TEXT_CHARS:
                break
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    lines = []
    # iter_inner_content keeps paragraphs and tables in document order.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    length = 0
    try:
        for sheet in workbook.worksheets:
            lines.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                if all(cell is None for cell in row):
                    continue
                line = "\t".join("" if cell is None else str(cell) for cell in row)
                lines.append(line)
                length += len(line)
                if length > MAX_TEXT_CHARS:
                    return "\n".join(lines)
    finally:
        workbook.close()
    return "\n".join(lines)


def _image(data: bytes) -> str:
    from PIL import Image

    return _ocr(Image.open(io.BytesIO(data)))


def _ocr(image) -> str:
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = _tesseract_cmd()
    try:
        return pytesseract.image_to_string(image, lang=get_settings().ocr_languages)
    except pytesseract.TesseractNotFoundError as exc:
        raise UnsupportedAttachmentError(
            "OCR needs the Tesseract program installed (set TESSERACT_CMD in .env)"
        ) from exc


@lru_cache
def _tesseract_cmd() -> str:
    configured = get_settings().tesseract_cmd
    if configured:
        return configured
    on_path = shutil.which("tesseract")
    if on_path:
        return on_path
    for root in ("LOCALAPPDATA", "ProgramFiles", "ProgramFiles(x86)"):
        base = os.environ.get(root)
        if not base:
            continue
        for sub in ("Programs/Tesseract-OCR", "Tesseract-OCR"):
            candidate = Path(base) / sub / "tesseract.exe"
            if candidate.exists():
                return str(candidate)
    return "tesseract"
